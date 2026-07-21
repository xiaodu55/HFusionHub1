"""
Word document parser (.docx)
"""

from typing import List
from app.core.parser.base import BaseParser, ParsedBlock, BlockType


class DocxParser(BaseParser):
    """Parse Word documents (.docx) into blocks"""

    def parse(self, file_path: str) -> List[ParsedBlock]:
        """Parse Word file into blocks"""
        try:
            from docx import Document
        except ImportError:
            raise ImportError("python-docx is required for Word parsing. Install with: pip install python-docx")

        doc = Document(file_path)
        blocks = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue

            # Determine block type based on style
            style_name = paragraph.style.name.lower()

            if 'heading' in style_name:
                # Extract heading level
                level = 1
                for i, char in enumerate(style_name):
                    if char.isdigit():
                        level = int(char)
                        break

                blocks.append(ParsedBlock(
                    content=text,
                    block_type=BlockType.HEADING,
                    level=level,
                    metadata={"style": paragraph.style.name}
                ))
            elif 'list' in style_name or paragraph.text.strip().startswith(('•', '-', '*', '1.', '2.', '3.')):
                blocks.append(ParsedBlock(
                    content=text,
                    block_type=BlockType.LIST,
                    metadata={"style": paragraph.style.name}
                ))
            else:
                blocks.append(ParsedBlock(
                    content=text,
                    block_type=BlockType.PARAGRAPH,
                    metadata={"style": paragraph.style.name}
                ))

        # Also extract tables
        for table in doc.tables:
            table_rows = []
            for row in table.rows:
                row_text = '| ' + ' | '.join(cell.text.strip() for cell in row.cells) + ' |'
                table_rows.append(row_text)

            if table_rows:
                content = '\n'.join(table_rows)
                blocks.append(ParsedBlock(
                    content=content,
                    block_type=BlockType.TABLE,
                    metadata={"rows": len(table.rows), "cols": len(table.columns)}
                ))

        return blocks
