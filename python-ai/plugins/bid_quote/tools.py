"""bid_quote — 报价表导出插件（HFusionHub 平台内建，容器沙箱执行）。

将投标报价明细（bid_quote items）渲染为 .xlsx 报价表，并内嵌评分点计算
（复用 bid_calc_scoring 的确定性逻辑，插件自包含、不依赖应用 ToolRegistry）。
容器契约见 docker/plugin-runner/run_tool.py：工具名经 PLUGIN_TOOL_NAME 传入，
JSON 输入经 PLUGIN_TOOL_INPUT 传入，函数以关键字参数调用，返回值须 JSON 可序列化。

工具：bid_export_quote
  输入（JSON 对象）：
    - items: [{name, spec, unit, qty, unit_price, tax_rate}]  报价明细（必填）
        name        品名/条目名（必填）
        spec        规格型号
        unit        单位（项/套/人天…）
        qty         数量（number）
        unit_price  单价（number，未税）
        tax_rate    税率（number，0.06 = 6%）
    - meta: {tender_number, project_title, bidder_name, deadline, currency}
    - points: [{name, max_score, weight?}]  评分点（可选，内嵌评分计算）
    - scores: [{name, score}]               实得分（可选）
 输出（JSON 对象）：
    - filename: str     建议文件名（含 .xlsx）
    - size_bytes: int   xlsx 字节数
    - currency: str
    - totals: {subtotal, tax_amount, grand_total}
    - lines: int        明细行数
    - scoring: {...}    评分汇总（提供了 points 时）
    - base64: str       xlsx 的 base64 编码（ASCII），调用方负责落盘/下发
"""

from __future__ import annotations

import base64
import io
import json
from typing import Any

PLUGIN_VERSION = "1.0.0"
PLUGIN_NAME = "bid_quote"


# ── 数值辅助（与 app/core/bid/tools.py 的 _to_number 保持一致）────────

def _to_number(value: Any) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _normalize_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _as_list(value: Any) -> list[dict[str, Any]]:
    if isinstance(value, str):
        try:
            value = json.loads(value)
        except json.JSONDecodeError:
            return []
    if not isinstance(value, list):
        return []
    return [v for v in value if isinstance(v, dict)]


# ── 评分点计算（复用 app/core/bid/tools.py::BidCalcScoringTool 逻辑）────

def _calc_scoring(points: list[dict[str, Any]], scores: list[dict[str, Any]]) -> dict[str, Any]:
    score_by_name: dict[str, float] = {}
    for item in scores:
        if not isinstance(item, dict) or "name" not in item:
            continue
        val = _to_number(item.get("score"))
        if val is not None:
            score_by_name[str(item["name"])] = val

    items: list[dict[str, Any]] = []
    total = 0.0
    max_total = 0.0
    weighted = 0.0
    weight_sum = 0.0
    for point in points:
        if not isinstance(point, dict):
            continue
        name = _normalize_text(point.get("name"))
        max_score = _to_number(point.get("max_score")) or 0.0
        weight = _to_number(point.get("weight")) or 0.0
        actual = score_by_name.get(name)
        if actual is None:
            items.append({
                "name": name,
                "max_score": max_score,
                "score": None,
                "percent": None,
                "note": "未提供该项得分",
            })
            max_total += max_score
            weight_sum += weight
            continue
        items.append({
            "name": name,
            "max_score": max_score,
            "score": actual,
            "percent": round(actual / max_score * 100, 2) if max_score else None,
            "weight": weight or None,
        })
        total += actual
        max_total += max_score
        weighted += actual * weight
        weight_sum += weight

    if not items:
        return {"error": "points 数组为空或不合法"}

    return {
        "items": items,
        "total_score": round(total, 2),
        "max_total": round(max_total, 2),
        "score_percent": round(total / max_total * 100, 2) if max_total else None,
        "weighted_total": round(weighted, 2) if weight_sum else None,
    }


# ── 报价表计算 ────────────────────────────────────────────────────────

def _calc_quote(items: list[dict[str, Any]]) -> dict[str, Any]:
    lines: list[dict[str, Any]] = []
    subtotal = 0.0
    tax_amount = 0.0
    for item in items:
        name = _normalize_text(item.get("name"))
        if not name:
            continue
        qty = _to_number(item.get("qty"))
        unit_price = _to_number(item.get("unit_price"))
        tax_rate = _to_number(item.get("tax_rate")) or 0.0
        if qty is None or unit_price is None:
            lines.append({
                "name": name,
                "spec": _normalize_text(item.get("spec")),
                "unit": _normalize_text(item.get("unit")),
                "qty": qty,
                "unit_price": unit_price,
                "tax_rate": tax_rate,
                "line_total": None,
                "note": "数量或单价缺失",
            })
            continue
        line_total = round(qty * unit_price, 2)
        line_tax = round(line_total * tax_rate, 2)
        subtotal += line_total
        tax_amount += line_tax
        lines.append({
            "name": name,
            "spec": _normalize_text(item.get("spec")),
            "unit": _normalize_text(item.get("unit")),
            "qty": qty,
            "unit_price": unit_price,
            "tax_rate": tax_rate,
            "line_total": line_total,
            "line_tax": line_tax,
        })
    if not lines:
        raise ValueError("items 必须是非空的报价明细数组 [{'name','qty','unit_price',...}]")
    return {
        "lines": lines,
        "totals": {
            "subtotal": round(subtotal, 2),
            "tax_amount": round(tax_amount, 2),
            "grand_total": round(subtotal + tax_amount, 2),
        },
    }


# ── 渲染 .xlsx ────────────────────────────────────────────────────────

def _render_xlsx(quote: dict[str, Any], meta: dict[str, Any], scoring: dict[str, Any] | None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # 报价明细
    ws = wb.active
    ws.title = "报价明细"
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    headers = ["品名", "规格", "单位", "数量", "单价(未税)", "税率", "金额(未税)", "税额"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for line in quote["lines"]:
        ws.append([
            line["name"],
            line.get("spec") or "",
            line.get("unit") or "",
            line.get("qty") if line.get("qty") is not None else "—",
            line.get("unit_price") if line.get("unit_price") is not None else "—",
            line.get("tax_rate"),
            line.get("line_total") if line.get("line_total") is not None else "—",
            line.get("line_tax") if line.get("line_tax") is not None else "—",
        ])
    t = quote["totals"]
    ws.append([])
    ws.append(["合计(未税)", "", "", "", "", "", t["subtotal"], t["tax_amount"]])
    ws.append(["含税总价", "", "", "", "", "", "", t["grand_total"]])
    for cell in ws[ws.max_row - 1]:
        cell.font = Font(bold=True)
    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 16

    # 评分汇总（可选）
    if scoring and "items" in scoring:
        ws2 = wb.create_sheet("评分汇总")
        ws2.append(["评分点", "满分", "实得分", "得分率%", "权重", "加权分"])
        for cell in ws2[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        for it in scoring["items"]:
            ws2.append([
                it.get("name"),
                it.get("max_score"),
                it.get("score") if it.get("score") is not None else "—",
                it.get("percent") if it.get("percent") is not None else "—",
                it.get("weight") if it.get("weight") is not None else "—",
                it.get("weight") * it.get("score")
                if it.get("score") is not None and it.get("weight")
                else "—",
            ])
        ws2.append([])
        ws2.append(["总分", scoring.get("max_total"), scoring.get("total_score"),
                    scoring.get("score_percent"), "", scoring.get("weighted_total")])
        for col in range(1, 7):
            ws2.column_dimensions[get_column_letter(col)].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def bid_export_quote(
    items: Any = None,
    meta: dict[str, Any] | None = None,
    points: Any = None,
    scores: Any = None,
    **kwargs: Any,
) -> dict[str, Any]:
    """把投标报价明细渲染为 .xlsx 报价表（可含评分汇总），返回 base64 字节与元信息。"""
    try:
        from openpyxl import Workbook  # noqa: F401 — 镜像内已安装，提前校验
    except ImportError as exc:  # pragma: no cover — 镜像内已安装
        raise RuntimeError(f"openpyxl 未安装: {exc}") from exc

    meta = meta if isinstance(meta, dict) else {}
    items = _as_list(items)
    if not items:
        raise ValueError("items 必须是非空的报价明细数组 [{'name','qty','unit_price',...}]")

    quote = _calc_quote(items)
    scoring = _calc_scoring(_as_list(points), _as_list(scores)) if points is not None else None

    raw = _render_xlsx(quote, meta, scoring)

    project_title = _normalize_text(meta.get("project_title")) or "投标报价"
    return {
        "filename": f"{project_title}-报价表.xlsx",
        "size_bytes": len(raw),
        "currency": _normalize_text(meta.get("currency")) or "¥",
        "totals": quote["totals"],
        "lines": len(quote["lines"]),
        "scoring": scoring,
        "base64": base64.b64encode(raw).decode("ascii"),
    }
